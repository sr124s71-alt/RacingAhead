import json, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
d=json.load(open('plan.json'))
D=lambda s: dt.date(*map(int,s.split('-')))
COL={'launch':'EB6834','design':'4A3AA7','identity':'2A78D6','notif':'1BAF7A','partner':'EDA100','booking':'E87BA4','ledger':'008300','web':'E34948','release':'8A8984','pm':'8A8984'}
H=Font(bold=True,color='FFFFFF'); HF=PatternFill('solid',fgColor='1F3864'); G=PatternFill('solid',fgColor='DCE3EF'); thin=Side(style='thin',color='C9CED8'); BD=Border(top=thin,bottom=thin,left=thin,right=thin)
wb=Workbook(); ws=wb.active; ws.title='Schedule'
ws['A1']='SportSeek Phase 2A — Integrated Project Plan (Baseline 0, 27 Sep 2026) — Srivin Platforms · Confidential'; ws['A1'].font=Font(bold=True,size=12,color='1F3864')
hdr=['WBS','Activity','Type','Owner','Start','Finish','Working days','Predecessors','Stream']
wk0=dt.date(2026,9,28); weeks=[wk0+dt.timedelta(weeks=i) for i in range(45)]
for i,h in enumerate(hdr+[w.strftime('%d-%b-%y') for w in weeks],1):
    c=ws.cell(3,i,h); c.font=H if i<=len(hdr) else Font(bold=True,color='FFFFFF',size=7); c.fill=HF; c.alignment=Alignment(wrap_text=True,vertical='center',textRotation=90 if i>len(hdr) else 0); c.border=BD
names={'S':'Summary','T':'Task','M':'Milestone'}; sname={'launch':'Release 2A.0 launch','design':'Discovery & design','identity':'Foundation & identity','notif':'Notifications','partner':'Partner & service','booking':'Booking & events','ledger':'Transaction ledger','web':'Web application','release':'Release, KT & warranty','pm':'Programme management'}
r=4
for x in d['rows']:
    vals=[x['id'],x['name'],names[x['type']],x['owner'],D(x['start']),D(x['end']),x['days'],x['pred'],sname[x['stream']]]
    for i,v in enumerate(vals,1):
        c=ws.cell(r,i,v); c.border=BD; c.alignment=Alignment(vertical='top',wrap_text=(i==2))
        if i in(5,6): c.number_format='dd-mmm-yy'
        if x['type']=='S': c.fill=G; c.font=Font(bold=True)
        if x['type']=='M': c.font=Font(bold=True)
    s,e=D(x['start']),D(x['end'])
    for j,w in enumerate(weeks):
        if s<=w+dt.timedelta(days=6) and e>=w:
            c=ws.cell(r,len(hdr)+1+j)
            if x['type']=='M': c.value='◆'; c.alignment=Alignment(horizontal='center'); c.font=Font(bold=True)
            else: c.fill=PatternFill('solid',fgColor='C5CCD8' if x['type']=='S' else COL[x['stream']])
    r+=1
for i,w in enumerate([7,62,10,10,11,11,8,20,20],1): ws.column_dimensions[get_column_letter(i)].width=w
for j in range(len(weeks)): ws.column_dimensions[get_column_letter(len(hdr)+1+j)].width=2.6
ws.row_dimensions[3].height=48; ws.freeze_panes='C4'; ws.auto_filter.ref=f'A3:{get_column_letter(len(hdr))}{r-1}'
ws2=wb.create_sheet('External dependencies')
h2=['ID','Dependency','Owner','Need-by','Unblocks (WBS)','Impact if late','Status','Comments']
for i,h in enumerate(h2,1): c=ws2.cell(1,i,h); c.font=H; c.fill=HF; c.border=BD
for k,x in enumerate(d['deps'],2):
    v=[x['id'],x['item'],x['owner'],D(x['needby']) if x['needby'][:2]=='20' else x['needby'],x['impacts'],x['ifLate'],'Open','']
    for i,val in enumerate(v,1):
        c=ws2.cell(k,i,val); c.border=BD; c.alignment=Alignment(wrap_text=True,vertical='top')
        if i==4 and not isinstance(val,str): c.number_format='dd-mmm-yy'
for i,w in enumerate([6,60,16,11,16,40,10,30],1): ws2.column_dimensions[get_column_letter(i)].width=w
ws2.freeze_panes='A2'
ws3=wb.create_sheet('Milestones')
for i,h in enumerate(['WBS','Milestone / gate','Owner','Date','Predecessors'],1): c=ws3.cell(1,i,h); c.font=H; c.fill=HF; c.border=BD
k=2
for x in d['rows']:
    if x['type']=='M':
        for i,val in enumerate([x['id'],x['name'],x['owner'],D(x['start']),x['pred']],1):
            c=ws3.cell(k,i,val); c.border=BD
            if i==4: c.number_format='dd-mmm-yy'
        k+=1
for i,w in enumerate([7,70,12,12,30],1): ws3.column_dimensions[get_column_letter(i)].width=w
ws4=wb.create_sheet('Read me')
for i,t in enumerate(['SportSeek Phase 2A — Project Plan workbook (SRV-SPS-2A-PLN-001a)','Companion to the WBS & Integrated Project Plan (SRV-SPS-2A-PLN-001). Baseline 0 dated 27 Sep 2026.',
 'Working days exclude weekends only; public/festival holidays are applied after the joint holiday calendar is agreed.','Predecessors: bare ID = Finish-to-Start; "SS" = Start-to-Start; "D-n" = external dependency (see External dependencies sheet).',
 'Gantt columns are weeks commencing Monday. Colours follow the workstream legend used in the plan document.'],1):
    ws4.cell(i,1,t).font=Font(bold=(i==1),size=12 if i==1 else 10)
ws4.column_dimensions['A'].width=130

rs=json.load(open('res.json'))
ws5=wb.create_sheet('Resource loading')
hh=['Skill group','Role','Seniority']+rs['months']+['Person-months','Key technical skills','Tools']
for i,h in enumerate(hh,1):
    c=ws5.cell(1,i,h); c.font=H; c.fill=HF; c.border=BD; c.alignment=Alignment(wrap_text=True,vertical='center')
for k,rr in enumerate(rs['roles'],2):
    vals=[rs['groups'][rr['group']],rr['role'],rr['seniority']]+rr['load']+[rr['pm'],rr['skills'],rr['tools']]
    for i,v in enumerate(vals,1):
        c=ws5.cell(k,i,v); c.border=BD; c.alignment=Alignment(wrap_text=True,vertical='top')
        if 4<=i<4+len(rs['months']) and v: c.fill=PatternFill('solid',fgColor=('9EC5F4' if v<1 else '5598E7' if v<2 else '2A78D6')); 
n=len(rs['roles'])+2
ws5.cell(n,2,'Total Srivin FTE').font=Font(bold=True)
for j in range(len(rs['months'])+1):
    col=get_column_letter(4+j); c=ws5.cell(n,4+j,f'=SUM({col}2:{col}{n-1})'); c.font=Font(bold=True); c.border=BD
for i,w in enumerate([22,34,18]+[8]*len(rs['months'])+[12,90,40],1): ws5.column_dimensions[get_column_letter(i)].width=w
ws5.freeze_panes='D2'

wb.move_sheet('Read me',offset=-4)
wb.save('out_plan.xlsx'); print('ok',r-4)
