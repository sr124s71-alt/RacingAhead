import json, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
g = json.load(open('gdr.json'))
H = Font(bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='1F3864'); th = Side(style='thin', color='C9CED8'); BD = Border(top=th, bottom=th, left=th, right=th)
SEV = {'Critical': 'F4C7C3', 'High': 'FCE5CD', 'Medium': 'FFF2CC', 'Low': 'E2EFDA'}
wb = Workbook(); wb.remove(wb.active)
def sheet(name, hdr, rows, widths, sevcol=None, datecols=(), status=None):
    ws = wb.create_sheet(name)
    for i, h in enumerate(hdr, 1):
        c = ws.cell(1, i, h); c.font = H; c.fill = HF; c.border = BD; c.alignment = Alignment(wrap_text=True, vertical='center')
    for r, row in enumerate(rows, 2):
        for i, v in enumerate(row, 1):
            if i in datecols and isinstance(v, str) and len(v) == 10 and v[4] == '-': v = dt.date.fromisoformat(v)
            c = ws.cell(r, i, v); c.border = BD; c.alignment = Alignment(wrap_text=True, vertical='top')
            if i in datecols and isinstance(v, dt.date): c.number_format = 'dd-mmm-yyyy'
        if sevcol and row[sevcol - 1] in SEV: ws.cell(r, sevcol).fill = PatternFill('solid', fgColor=SEV[row[sevcol - 1]])
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'B2'; ws.auto_filter.ref = f'A1:{get_column_letter(len(hdr))}{len(rows) + 1}'
    if status:
        dv = DataValidation(type='list', formula1='"' + ','.join(status) + '"', allow_blank=True); ws.add_data_validation(dv)
        col = get_column_letter(hdr.index('Status') + 1); dv.add(f'{col}2:{col}{len(rows) + 200}')
    return ws
ST = ['Open', 'Open — urgent', 'Decided', 'Closed', 'Tracked']
rm = wb.create_sheet('Read me')
for i, t in enumerate(['SportSeek Phase 2A — Gap & Decision Register tracker (SRV-SPS-2A-GDR-004a)', 'Companion to the Gap & Decision Register v1.0, 15 October 2026. Prepared by Srivin Platforms.',
    'C-n = conflict between SportSeek documents; G-n = gap in Srivin architecture; DEC-n = decision. Every C and G links to one DEC.',
    'Update Status, Decision date, Outcome and Decided by in the Decisions sheet as decisions are taken. Status values: Open, Open — urgent, Decided, Closed, Tracked.',
    'Regulatory references are flagged for architecture impact only and need confirmation by SportSeek\'s legal and tax advisers.'], 1):
    rm.cell(i, 1, t).font = Font(bold=(i == 1), size=12 if i == 1 else 10)
rm.column_dimensions['A'].width = 140
sheet('Conflicts', ['ID', 'Title', 'Severity', 'Vision says', 'SOW says', 'Impact', 'Options', 'Srivin recommendation', 'Decision', 'Status'],
      [[c['id'], c['title'], c['sev'], c['vision'], c['sow'], c['impact'], c['options'], c['rec'], c['dec'], 'Open'] for c in g['conflicts']],
      [6, 28, 10, 45, 45, 45, 50, 50, 9, 12], sevcol=3, status=ST)
sheet('Gaps', ['ID', 'Gap', 'Severity', 'Evidence (source)', 'Impact', 'Recommended resolution', 'When', 'Owner', 'Decision', 'Status'],
      [[x['id'], x['title'], x['sev'], x['evidence'], x['impact'], x['rec'], x['when'], x['owner'], x['dec'], 'Open'] for x in g['gaps']],
      [6, 32, 10, 45, 45, 50, 24, 20, 9, 12], sevcol=3, status=ST)
ws = sheet('Decisions', ['ID', 'Decision required', 'Srivin recommendation', 'Owner', 'Need-by', 'If late', 'Status', 'Decision date', 'Outcome', 'Decided by'],
      [list(d) + ['', '', ''] for d in sorted(g['decisions'], key=lambda d: d[4] if d[4][:2] == '20' else '9999')],
      [9, 36, 50, 24, 13, 36, 14, 13, 40, 18], datecols=(5, 8), status=ST)
ws.conditional_formatting.add('G2:G200', CellIsRule(operator='equal', formula=['"Open — urgent"'], fill=PatternFill('solid', fgColor='F4C7C3')))
sheet('Traceability', ['Vision area', 'Capability (vision boards)', 'Architecture component / phase', 'Status', 'Register items'], [list(t) for t in g['trace']], [16, 55, 45, 16, 16])
sheet('Best practices', ['Area', 'Practice', 'Standard / reference', 'How it applies to SportSeek', 'Status', 'Link'], [list(p) for p in g['practices']], [14, 45, 40, 60, 16, 14])
wb.save('out_gdr.xlsx'); print('ok')
